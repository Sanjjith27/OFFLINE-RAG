import os
import json

import pandas as pd
import fitz

from docx import Document
from pptx import Presentation

import shutil
from PIL import Image
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


# =========================================================
# TESSERACT CONFIGURATION
# =========================================================

def find_tesseract():
    """Detect Tesseract OCR executable on Windows or system PATH."""
    env_path = os.environ.get("TESSERACT_PATH")
    if env_path and os.path.exists(env_path):
        return env_path

    which_path = shutil.which("tesseract")
    if which_path:
        return which_path

    common_paths = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tesseract.exe"),
        r"C:\tools\tesseract\tesseract.exe",
    ]
    for p in common_paths:
        if p and os.path.exists(p):
            return p

    return None


# Initialize tesseract_cmd if found
_tesseract_bin = find_tesseract()
if _tesseract_bin:
    pytesseract.pytesseract.tesseract_cmd = _tesseract_bin


# =========================================================
# PDF LOADER
# =========================================================

def load_pdf(path):
    doc = fitz.open(path)
    documents = []

    for page_number, page in enumerate(doc):
        text = page.get_text().strip()

        if text:
            documents.append({
                "text": text,
                "source": os.path.basename(path),
                "page": page_number + 1,
                "type": "pdf"
            })

    return documents


# =========================================================
# DOCX LOADER
# =========================================================

def load_docx(path):
    doc = Document(path)

    text = "\n".join(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    )

    return [{
        "text": text,
        "source": os.path.basename(path),
        "page": None,
        "type": "docx"
    }]


# =========================================================
# PPTX LOADER
# =========================================================

def load_pptx(path):
    presentation = Presentation(path)
    documents = []

    for slide_number, slide in enumerate(presentation.slides):
        text = ""

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

        if text.strip():
            documents.append({
                "text": text.strip(),
                "source": os.path.basename(path),
                "page": slide_number + 1,
                "type": "pptx"
            })

    return documents


# =========================================================
# TXT LOADER
# =========================================================

def load_txt(path):
    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:
        text = file.read()

    return [{
        "text": text,
        "source": os.path.basename(path),
        "page": None,
        "type": "txt"
    }]


# =========================================================
# CSV LOADER
# =========================================================

def load_csv(path):
    df = pd.read_csv(path)

    return [{
        "text": df.to_string(index=False),
        "source": os.path.basename(path),
        "page": None,
        "type": "csv"
    }]


# =========================================================
# XLSX / XLS LOADER
# =========================================================

def load_xlsx(path):
    excel = pd.ExcelFile(path)
    documents = []

    for sheet in excel.sheet_names:
        df = pd.read_excel(
            path,
            sheet_name=sheet
        )

        documents.append({
            "text": df.to_string(index=False),
            "source": os.path.basename(path),
            "page": sheet,
            "type": "xlsx"
        })

    return documents


# =========================================================
# JSON LOADER
# =========================================================

def load_json(path):
    with open(
        path,
        "r",
        encoding="utf-8"
    ) as file:
        data = json.load(file)

    return [{
        "text": json.dumps(
            data,
            indent=2
        ),
        "source": os.path.basename(path),
        "page": None,
        "type": "json"
    }]


# =========================================================
# IMAGE OCR LOADER
# =========================================================

def load_image(path):
    try:
        # Dynamic check in case Tesseract was installed while server was running
        tesseract_bin = find_tesseract()
        if tesseract_bin:
            pytesseract.pytesseract.tesseract_cmd = tesseract_bin
        else:
            return [{
                "text": "OCR engine (Tesseract) is not installed or not found. Please install Tesseract OCR (e.g. run tesseract-installer.exe in the project root) to enable image text extraction.",
                "source": os.path.basename(path),
                "page": None,
                "type": "image"
            }]

        image = Image.open(path)

        # Convert image to RGB
        image = image.convert("RGB")

        # Extract text using Tesseract OCR
        text = pytesseract.image_to_string(
            image
        ).strip()

        if not text:
            text = "No readable text was found in this image."

        return [{
            "text": text,
            "source": os.path.basename(path),
            "page": None,
            "type": "image"
        }]

    except Exception as e:
        return [{
            "text": f"OCR failed for this image: {str(e)}",
            "source": os.path.basename(path),
            "page": None,
            "type": "image"
        }]


# =========================================================
# MAIN DOCUMENT LOADER
# =========================================================

def load_document(path):

    extension = os.path.splitext(
        path
    )[1].lower()

    if extension == ".pdf":
        return load_pdf(path)

    elif extension == ".docx":
        return load_docx(path)

    elif extension == ".pptx":
        return load_pptx(path)

    elif extension == ".txt":
        return load_txt(path)

    elif extension == ".csv":
        return load_csv(path)

    elif extension in [".xlsx", ".xls"]:
        return load_xlsx(path)

    elif extension == ".json":
        return load_json(path)

    elif extension in [
        ".png",
        ".jpg",
        ".jpeg"
    ]:
        return load_image(path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )