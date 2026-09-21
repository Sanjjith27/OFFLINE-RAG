import os
import json
import pandas as pd
import fitz

from docx import Document
from pptx import Presentation


def load_pdf(path):

    doc = fitz.open(path)

    text = ""

    for page_number, page in enumerate(doc):

        page_text = page.get_text()

        text += (
            f"\n[Page {page_number + 1}]\n"
            f"{page_text}"
        )

    return text


def load_docx(path):

    doc = Document(path)

    text = ""

    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"

    return text


def load_pptx(path):

    presentation = Presentation(path)

    text = ""

    for slide_number, slide in enumerate(
        presentation.slides
    ):

        text += f"\n[Slide {slide_number + 1}]\n"

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


def load_txt(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return file.read()


def load_csv(path):

    df = pd.read_csv(path)

    return df.to_string(index=False)


def load_xlsx(path):

    excel = pd.ExcelFile(path)

    text = ""

    for sheet in excel.sheet_names:

        df = pd.read_excel(
            path,
            sheet_name=sheet
        )

        text += f"\n[Sheet: {sheet}]\n"
        text += df.to_string(index=False)

    return text


def load_json(path):

    with open(
        path,
        "r",
        encoding="utf-8"
    ) as file:

        data = json.load(file)

    return json.dumps(
        data,
        indent=2
    )


def load_document(path):

    extension = os.path.splitext(path)[1].lower()

    if extension == ".pdf":
        return load_pdf(path)

    elif extension == ".docx":
        return load_docx(path)

    elif extension == ".pptx":
        return load_pptx(path)

    elif extension == ".txt":
        return load_txt(path)

    elif extension == ".csv":
        return load_csv(path)

    elif extension in [".xlsx", ".xls"]:
        return load_xlsx(path)

    elif extension == ".json":
        return load_json(path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )